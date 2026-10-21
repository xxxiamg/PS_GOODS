import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
import app_settings_config as cfg


def format_data_structures(fin_A, fin_B, fin_C, hipo):
    def build_sheet(data, t):
        r_lbl, s_lbl, raw_st = [], [], data["Остаток"].copy()
        data["Остаток"] = data["Остаток"].astype(str)

        for i in range(len(data)):
            p, st_num = i + 1, int(raw_st.at[i])
            sold_num = int(data.at[i, "P"])

            if t == "A":
                if p <= 5:
                    r_lbl.append("TOP 5")
                elif p <= 10:
                    r_lbl.append("TOP 10")
                elif data.at[i, "Товар"] in hipo["Товар"].values:
                    r_lbl.append("HiPo")
                else:
                    r_lbl.append("Stable stock")

                if st_num == 0:
                    s_lbl.append("🔥 СРОЧНЫЙ ДОЗАКАЗ")
                    data.at[i, "Остаток"] = "Нет на складе"
                elif st_num >= 4 and sold_num <= 3:
                    s_lbl.append("⚠️ ИЗБЫТОК СТОКА")
                elif st_num >= 3:
                    s_lbl.append("По плану")
                else:
                    s_lbl.append("Пополнение стока")

            elif t == "B":
                s_lbl.append("⚠️ Только под заказ")
                if st_num == 0:
                    data.at[i, "Остаток"] = "Нет на складе"
            else:
                s_lbl.append("⛔️ СТОП-ЗАКУПКА")
                if st_num == 0:
                    data.at[i, "Остаток"] = "Нет на складе"

        data.insert(4, "Статус закупки", s_lbl)
        data.rename(columns={"Товар": "Наименование товара",
                    "P": "Продано", "C": "Цикл продаж"}, inplace=True)

        if t == "A":
            data.insert(0, "Рейтинг", r_lbl)
            data["Оборот"] = data["Продано"] + \
                data["Остаток"].replace("Нет на складе", 0).astype(int)
            return data[["Рейтинг", "Наименование товара", "Продано", "Остаток", "Статус закупки", "Дней на складе", "Оборот", "Цикл продаж", "O"]]
        else:
            return data[["Наименование товара", "Продано", "Остаток", "Статус закупки", "Дней на складе", "Цикл продаж", "O"]]

    df_A = build_sheet(fin_A, "A")
    df_B = build_sheet(fin_B, "B")
    df_C = build_sheet(fin_C, "C")

    t5_e = df_A[df_A["Рейтинг"] == "TOP 5"].index.max(
    ) if not df_A[df_A["Рейтинг"] == "TOP 5"].empty else 0
    t10_e = df_A[df_A["Рейтинг"] == "TOP 10"].index.max(
    ) if not df_A[df_A["Рейтинг"] == "TOP 10"].empty else 0
    hipo_e = df_A[df_A["Рейтинг"] == "HiPo"].index.max(
    ) if not df_A[df_A["Рейтинг"] == "HiPo"].empty else 0

    emp = pd.Series([None] * len(df_A.columns),
                    index=df_A.columns).to_frame().T

    # Разрезаем таблицу на 4 части и склеиваем с двойными пустыми строками во всех точках перехода
    df_A_spaced = pd.concat([
        df_A.loc[:t5_e],
        emp, emp,
        df_A.loc[t5_e+1:t10_e],
        emp, emp,
        df_A.loc[t10_e+1:hipo_e],
        emp, emp,
        df_A.loc[hipo_e+1:]
    ]).reset_index(drop=True)

    return df_A_spaced, df_B, df_C


def apply_excel_styles(workbook):
    f_low = PatternFill("solid", fgColor=cfg.COLOR_LOW)
    f_med = PatternFill("solid", fgColor=cfg.COLOR_MED)
    f_high = PatternFill("solid", fgColor=cfg.COLOR_HIGH)
    f_purp = PatternFill("solid", fgColor=cfg.COLOR_PURPLE)
    f_amb = PatternFill("solid", fgColor=cfg.COLOR_HEADER)
    f_t510 = PatternFill("solid", fgColor=cfg.COLOR_TOP)
    f_hp = PatternFill("solid", fgColor=cfg.COLOR_HIPO)

    for name in workbook.sheetnames:
        ws = workbook[name]
        ws.row_dimensions.height = 22.5

        for r_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[r_idx].height = 15

        for c_idx in range(1, ws.max_column + 1):
            col_name = ws.cell(row=1, column=c_idx).value
            if col_name in cfg.COLUMN_WIDTHS:
                ws.column_dimensions[ws.cell(
                    row=1, column=c_idx).column_letter].width = cfg.COLUMN_WIDTHS[col_name]

        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=1):
            for cell in row:
                cell.font = Font(bold=(row_idx == 1))
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True)
                if row_idx == 1:
                    cell.fill = f_amb

        c_tr, c_o, c_p, c_st = None, None, None, None
        for c in range(1, ws.max_column + 1):
            val = ws.cell(1, c).value
            if val == "Оборот":
                c_tr = c
            if val == "O":
                c_o = c
            if val == "Продано":
                c_p = c
            if val == "Остаток":
                c_st = c

        for r in range(2, ws.max_row + 1):
            if name == "Группа А" and ws.cell(r, 2).value is None:
                continue

            o_val = ws.cell(r, c_o).value

            if c_tr:
                turn_cell = ws.cell(r, c_tr)
                turn_cell.value = str(o_val).split("-")[-1]
                if "низкий" in str(o_val):
                    turn_cell.fill = f_low
                elif "средний" in str(o_val):
                    turn_cell.fill = f_med
                elif "высокий" in str(o_val):
                    turn_cell.fill = f_high

            if name == "Группа А":
                rc = ws.cell(r, 1)
                rt = str(rc.value)

                if "TOP 5" in rt or "TOP 10" in rt:
                    rc.fill = f_t510
                    ws.row_dimensions[r].height = 15
                elif "HiPo" in rt or "Stable Stock" in rt:
                    rc.fill = f_hp if "HiPo" in rt else PatternFill(
                        fill_type=None)
                    ws.row_dimensions[r].height = 30

                if ws.cell(r, 5).value == "🔥 СРОЧНЫЙ ДОЗАКАЗ":
                    ws.row_dimensions[r].height = 22.5

            if name == "Группа B" and ws.cell(r, c_p).value == 2 and ws.cell(r, c_st).value in [0, "0", "Нет на складе"]:
                for c_idx in range(1, ws.max_column):
                    ws.cell(r, c_idx).fill = f_purp

        ws.delete_cols(c_o)
